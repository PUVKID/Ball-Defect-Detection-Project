import pathlib,json
from copy import deepcopy
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
C=json.loads(pathlib.Path('work/content_v2.json').read_text(encoding='utf8'));d=Document()
for st in d.styles:
 if hasattr(st,'font'):st.font.bold=False
 if st.element.pPr is not None:
  for b in list(st.element.pPr.findall(qn('w:pBdr'))):st.element.pPr.remove(b)
d.core_properties.author='';d.core_properties.title=C['title'].replace('\n',' ')
fn={'body':'휴먼명조','heading':'휴먼고딕','caption':'HY중고딕','engtitle':'HY견명조','english':'Times New Roman','abstract_heading':'돋움'}
def fmt_run(r,f,size,spacing=-.06):
 r.font.bold=False;r.font.name=fn[f];r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0)
 rp=r._element.get_or_add_rPr();rf=rp.rFonts
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),fn[f])
 w=OxmlElement('w:w');w.set(qn('w:val'),'90');rp.append(w)
 s=OxmlElement('w:spacing');s.set(qn('w:val'),str(round(size*spacing*20)));rp.append(s)
 return r
def p(s='',f='body',size=10,leading=14,align='justify',before=0,after=5,indent=9,style=None,space=-.06):
 x=d.add_paragraph(style=style);fmt_run(x.add_run(s),f,size,space)
 pf=x.paragraph_format;pf.line_spacing=Pt(leading);pf.space_before=Pt(before);pf.space_after=Pt(after);pf.first_line_indent=Pt(indent);pf.widow_control=False
 x.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY,'left':WD_ALIGN_PARAGRAPH.LEFT}[align]
 return x
def section(kind=None,cols=2):
 s=d.sections[0] if kind is None else d.add_section(kind)
 s.page_width=Mm(190);s.page_height=Mm(260);s.left_margin=s.right_margin=Mm(20);s.top_margin=s.bottom_margin=Mm(17.5);s.header_distance=Mm(7);s.footer_distance=Mm(6)
 co=s._sectPr.find(qn('w:cols'));co.set(qn('w:num'),str(cols));co.set(qn('w:space'),str(round(8/25.4*1440)))
 return s
s=section(cols=1)
hp=s.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fmt_run(hp.add_run('- 2026 한국정보기술학회 추계 종합학술대회 논문집 -'),'abstract_heading',10,spacing=0);hp.paragraph_format.space_after=Pt(0)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;fmt_run(f.add_run('- '),'english',10)
r=f.add_run();fmt_run(r,'english',10);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fr=OxmlElement('w:r');fr.append(deepcopy(r._r.rPr));ft=OxmlElement('w:t');ft.text='1';fr.append(ft);fld.append(fr);r._r.addnext(fld);fmt_run(f.add_run(' -'),'english',10)
p(C['title'],'body',17,21,'center',after=7,indent=0,style='Title',space=-.07)
p('','body',11,11,'center',after=2,indent=0)
p(C['english'],'engtitle',15,19,'center',after=7,indent=0,space=-.07)
p('','body',10,10,'center',after=4,indent=0)
section(WD_SECTION.CONTINUOUS,2)
for i,items in enumerate(C['columns']):
 if i in [2,4]:section(WD_SECTION.NEW_PAGE,2)
 elif i%2:
  x=p('',leading=1,after=0,indent=0);x.add_run().add_break(WD_BREAK.COLUMN)
 for kind,text in items:
  if kind=='h':p(text,'abstract_heading' if text in ['요  약','Abstract'] else 'heading',10.5 if text in ['요  약','Abstract'] else 11,15,'center',before=2,after=6,indent=0).paragraph_format.keep_with_next=True
  elif kind=='sub':p(text,'caption',11,15,'left',before=2,after=5,indent=0).paragraph_format.keep_with_next=True
  elif kind=='ref':p(text,'english',10,11,'left',after=4,indent=0)
  elif kind=='abs':p(text,'body',9.2,13.2,after=7)
  elif kind=='en':p(text,'english',9.2,12.7,after=7)
  elif kind=='kw':p(text,'english',9.2,12,'left',after=7,indent=0)
  elif kind=='fig':
   x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(3);x.paragraph_format.line_spacing=1;x.paragraph_format.keep_with_next=True
   x.add_run().add_picture(f'work/restyled_horizontal/paper_figure{text}.png',width=Mm(71))
  elif kind=='cap':p(text,'caption',9,11,'center',after=7,indent=0)
  else:p(text)
for x in d.paragraphs:
 if x._p.pPr is not None and x._p.pPr.find(qn('w:sectPr')) is not None:
  x.paragraph_format.space_after=Pt(0);x.paragraph_format.space_before=Pt(0);x.paragraph_format.line_spacing=Pt(1);fmt_run(x.add_run(''),'body',1)
d.save('work/paper_v2.docx')
print('Wrote revised two-column paper')
